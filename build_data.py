"""Build the market dashboard data files from universe.csv.

Outputs (in ./data):
  master_data_file.csv  daily series: a metadata block (type, name, code, frequency, field,
                        currency, unit, source), a blank row, then one row per date and one
                        column per instrument. Weekly/monthly/quarterly/annual series go to
                        master_data_file_weekly/_monthly/_quarterly/_annual.csv in the same layout.
  master_data_long.csv  the same data in long format (one row per date x instrument x field)
  master_dashboard.csv  the long format with only date, frequency, code, field and value: what the
                        dashboard reads (its data files must stay under 8 MB)
  market_snapshot.csv   one row per instrument: latest value, 1D-5Y changes, 50/200-day averages,
                        52-week range, 20-day volatility and signals (see daily_stats)
  Market_Data.xlsx      Snapshot, Master (one sheet per frequency, same layout) and Universe sheets
  earnings_calendar.csv earnings dates and consensus for Equity Stock rows (see build_calendar.py)
The files in ./data hold the last YEARS of history; the full history is in the archive below.

Archive: archive/market_archive.parquet, one row per date x instrument x field (date, frequency,
code, field, value, source, fetched = when the value was last fetched). Nothing is ever trimmed.
Each run fetches a series only from OVERLAP days before its last archived date and merges that in;
a new instrument gets history from HISTORY_START. If the re-fetched days differ from the archive (forward-adjusted
prices after a dividend or split, or a revised value), the series' whole history is fetched again.
The archive before the run is kept as market_archive.prev.parquet. The first run seeds the archive
from master_data_long.csv. Instruments removed from universe.csv stay in the archive.

  python build_data.py              normal run
  python build_data.py --backfill   also fetch older history so every series reaches back to HISTORY_START

universe.csv columns:
  code       Wind code (or Yahoo ticker when source = yahoo)
  type       grouping shown on the dashboard, e.g. Broad Market Index, Equity Stock, Macro
  subtype    optional sub-heading within a type, e.g. Global Markets, APAC Markets; cards follow file order
  country    Equity Stock only: where the company is run from (head office), e.g. United States, Taiwan;
             groups the calendar's stock filter (group order: COUNTRY_ORDER in build_calendar.py)
  domicile   Equity Stock only: country of incorporation (Wind wss field address), e.g. Cayman Islands for Fabrinet
  frequency  D daily, W weekly, M monthly, Q quarterly, A annual
  field      Wind field to pull, e.g. close
  calendar   exchange trading calendar for Wind (NYSE, NASDAQ, CBOE, HKEX, SSE, TSE, KSE, ASX, LSE, XETRA, EPA,
             ASE, B3, BMV...); blank = Wind default (SSE); OWN = the series' own trading days, for markets
             Wind has no calendar for (Italy, Spain) and MSCI indices: weekdays, dropping days with no daily change
  adjust     Wind PriceAdj: F forward-adjusted (stocks, ETFs), blank = none (indices)
  source     wind (WSD time series), wind_edb (macro series), yahoo, eia (EIA daily spot price page;
             code = EIA series id, e.g. EER_EPJK_PF4_RGC_DPG for US Gulf Coast jet fuel)
  headline   Y to show as a headline tile on the dashboard
  chg_type   pct (percent change) or bp (basis-point change, for yields)
  earnings_session  Equity Stock only: AMC (reports after the US close), BMO (before the open), blank = unknown

Needs Wind Terminal open and logged in for Wind sources. If one instrument fails,
its archived history is kept and the failure is printed.
"""
import csv
import json
import math
import os
import re
import shutil
import statistics
import sys
import urllib.parse
import urllib.request
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "data"
LONG = OUT / "master_data_long.csv"
DASH = OUT / "master_dashboard.csv"
DASH_COLS = ["date", "frequency", "code", "field", "value"]
ARCHIVE = ROOT / "archive" / "market_archive.parquet"
ARCHIVE_SCHEMA = pa.schema([
    ("date", pa.date32()), ("frequency", pa.string()), ("code", pa.string()), ("field", pa.string()),
    ("value", pa.float64()), ("source", pa.string()), ("fetched", pa.date32()),
])
YEARS = 5  # history written to ./data
HISTORY_START = date(2020, 1, 1)  # history fetched for a new instrument, and by --backfill
OVERLAP = {"D": 30, "W": 70, "M": 190, "Q": 400, "A": 800}  # days re-fetched before the last archived date
SGT = timezone(timedelta(hours=8))
WIND_PERIOD = {"D": "D", "W": "W", "M": "M", "Q": "Q", "A": "Y"}
MASTER_COLS = ["date", "frequency", "code", "name", "type", "field", "value", "unit", "currency", "source"]
META = ["type", "name", "code", "frequency", "field", "currency", "unit", "source"]
WIDE_FILES = {
    "D": ("master_data_file.csv", "Master"),
    "W": ("master_data_file_weekly.csv", "Master_W"),
    "M": ("master_data_file_monthly.csv", "Master_M"),
    "Q": ("master_data_file_quarterly.csv", "Master_Q"),
    "A": ("master_data_file_annual.csv", "Master_A"),
}
HORIZONS = ["chg_1d", "chg_1w", "chg_1m", "chg_3m", "chg_6m", "chg_ytd", "chg_1y", "chg_2y", "chg_3y", "chg_5y"]
STATS = ["ma50", "ma200", "vs_ma50", "vs_ma200", "high_52w", "low_52w", "from_high_52w", "pos_52w",
         "vol_20d", "sigma_1d", "signal_52w", "cross_200"]


def years_ago(d, n):
    try:
        return d.replace(year=d.year - n)
    except ValueError:  # 29 Feb
        return d.replace(year=d.year - n, day=28)


def as_date(t):
    return t.date() if isinstance(t, datetime) else t


_wind = None


def wind():
    global _wind
    if _wind is None:
        from WindPy import w
        r = w.start(waitTime=60)
        if r.ErrorCode != 0:
            raise RuntimeError(f"Wind did not start (error {r.ErrorCode}): open Wind Terminal and log in")
        _wind = w
    return _wind


def fetch_wind(u, begin, end):
    own = u["calendar"].strip().upper() == "OWN" and u["frequency"] == "D"
    opts = [f"Period={WIND_PERIOD[u['frequency']]}"]
    if own:
        opts += ["Days=Weekdays", "Fill=Blank"]
    elif u["calendar"]:
        opts.append(f"TradingCalendar={u['calendar']}")
    if u["adjust"]:
        opts.append(f"PriceAdj={u['adjust']}")
    fields = u["field"] + (",pct_chg" if own else "")
    # OWN starts 10 days early so the first requested day has a previous close to compare with
    start = begin - timedelta(days=10) if own else begin
    d = wind().wsd(u["code"], fields, start.isoformat(), end.isoformat(), ";".join(opts))
    if d.ErrorCode != 0:
        raise RuntimeError(f"Wind error {d.ErrorCode}: {d.Data}")
    if not own:
        # NaN before a stock listed (SpaceX, CoreWeave, IREN within the 5-year window)
        return [(as_date(t), v) for t, v in zip(d.Times, d.Data[0]) if v is not None and v == v]
    # Wind fills the close on a day the market was shut but leaves its daily change blank,
    # or 0 with the same close when the index was not calculated (e.g. MSCI on 25 Dec)
    out = []
    for t, v, chg in zip(d.Times, d.Data[0], d.Data[1]):
        if v is None or v != v or chg is None or chg != chg:
            continue
        if chg == 0 and out and out[-1][1] == v:
            continue
        out.append((as_date(t), v))
    return [(t, v) for t, v in out if t >= begin]


def fetch_wind_edb(u, begin, end):
    d = wind().edb(u["code"], begin.isoformat(), end.isoformat(), "Fill=Blank")
    if d.ErrorCode != 0:
        raise RuntimeError(f"Wind EDB error {d.ErrorCode}: {d.Data}")
    return [(as_date(t), v) for t, v in zip(d.Times, d.Data[0]) if v is not None]


def fetch_yahoo(u, begin, end):
    start = int(datetime.combine(begin, datetime.min.time(), timezone.utc).timestamp())
    stop = int(datetime.now(timezone.utc).timestamp())
    url = (
        "https://query1.finance.yahoo.com/v8/finance/chart/"
        + urllib.parse.quote(u["code"])
        + f"?period1={start}&period2={stop}&interval=1d"
    )
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=30) as r:
        res = json.load(r)["chart"]["result"][0]
    offset = res["meta"].get("gmtoffset", 0)
    series = {}
    for ts, c in zip(res.get("timestamp", []), res["indicators"]["quote"][0]["close"]):
        if c is not None:
            series[datetime.fromtimestamp(ts + offset, timezone.utc).date()] = c
    return sorted(series.items())


def fetch_eia(u, begin, end):
    # The daily history page has one row per week ("2026 Oct-19 to Oct-23") with Monday-Friday cells
    if u["frequency"] != "D":
        raise RuntimeError("eia source supports daily series only")
    url = f"https://www.eia.gov/dnav/pet/hist/{u['code'].lower()}D.htm"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=60) as r:
        page = r.read().decode("latin-1")
    out = []
    for label, cells in re.findall(r"<tr>\s*<td class='B6'>(.*?)</td>(.*?)</tr>", page, re.S):
        monday = datetime.strptime(label.replace("&nbsp;", " ").split(" to ")[0].replace(" ", ""), "%Y%b-%d").date()
        for i, v in enumerate(re.findall(r"<td class='B3'>([^<]*)</td>", cells)):
            d = monday + timedelta(days=i)
            if v.strip() and begin <= d <= end:
                out.append((d, float(v)))
    return out


FETCH = {"wind": fetch_wind, "wind_edb": fetch_wind_edb, "yahoo": fetch_yahoo, "eia": fetch_eia}
SOURCE_LABEL = {"wind": "Wind", "wind_edb": "Wind EDB", "yahoo": "Yahoo", "eia": "EIA"}


def load_archive():
    """{(code, field, frequency): {date: (value, source, fetched)}}"""
    arc = defaultdict(dict)
    if ARCHIVE.exists():
        t = pq.read_table(ARCHIVE).to_pydict()
        for d, freq, code, field, v, src, fetched in zip(*(t[c] for c in ARCHIVE_SCHEMA.names)):
            arc[(code, field, freq)][d] = (v, src, fetched)
    elif LONG.exists():
        seeded = date.fromtimestamp(LONG.stat().st_mtime)
        for r in read_csv(LONG):
            arc[(r["code"], r["field"], r["frequency"])][date.fromisoformat(r["date"])] = (float(r["value"]), r["source"], seeded)
        print(f"Archive started from {LONG.name}")
    return arc


def save_archive(arc):
    # Written to a temporary file first, so a failed write never replaces the good archive
    cols = {c: [] for c in ARCHIVE_SCHEMA.names}
    for (code, field, freq), rows in sorted(arc.items()):
        for d in sorted(rows):
            v, src, fetched = rows[d]
            for c, x in zip(ARCHIVE_SCHEMA.names, (d, freq, code, field, v, src, fetched)):
                cols[c].append(x)
    ARCHIVE.parent.mkdir(exist_ok=True)
    tmp = ARCHIVE.with_name(ARCHIVE.name + ".tmp")
    pq.write_table(pa.table(cols, schema=ARCHIVE_SCHEMA), tmp, compression="zstd")
    if ARCHIVE.exists():
        shutil.copyfile(ARCHIVE, ARCHIVE.with_name(ARCHIVE.stem + ".prev.parquet"))
    os.replace(tmp, ARCHIVE)


def differs(a, b):
    return abs(a - b) > max(1.5e-4, 1e-6 * abs(b))


def refresh(u, have, today, backfill):
    """Fetch what the archive lacks for one instrument and merge it into `have` ({date: (value, source, fetched)}).
    Returns a note for the log ("" for a routine update)."""
    label = SOURCE_LABEL.get(u["source"], u["source"])

    def get(begin, end=today):
        return sorted(FETCH[u["source"].strip().lower()](u, begin, end))

    def merge(new, start):
        # The source's answer replaces everything archived from `start` on, so days it no longer has go too
        for d in [d for d in have if d >= start]:
            del have[d]
        have.update({d: (round(v, 6), label, today) for d, v in new})

    notes = []
    if not have:
        begin = HISTORY_START
        new = get(begin)
        if not new:
            raise RuntimeError("no data returned")
        merge(new, begin)
        notes.append(f"new series, {len(new)} rows from {new[0][0]}")
    else:
        begin = max(have) - timedelta(days=OVERLAP[u["frequency"]])
        new = get(begin)
        if not new:
            raise RuntimeError("no data returned")
        changed = [d for d, v in new if d in have and differs(have[d][0], v)]
        if not changed:
            merge(new, begin)
        else:
            # Earlier values changed (forward-adjusted prices after a dividend or split, or a revision)
            first = min(have)
            new = get(first)
            if not new:
                raise RuntimeError("history changed but the re-fetch returned no data")
            merge(new, new[0][0])
            notes.append(f"history changed (seen from {changed[0]}); re-fetched {len(new)} rows from {new[0][0]}")
            if new[0][0] - first > timedelta(days=14):
                notes.append(f"rows before {new[0][0]} were not returned again and stay as archived")
    if backfill:
        first = min(have)
        if first - HISTORY_START > timedelta(days=14):
            try:
                old = [(d, v) for d, v in get(HISTORY_START, first - timedelta(days=1)) if d < first]
                have.update({d: (round(v, 6), label, today) for d, v in old})
                notes.append(f"backfilled {len(old)} rows" + (f" from {old[0][0]}" if old else " (none earlier)"))
            except Exception as e:
                notes.append(f"backfill failed: {e}")
    return "; ".join(notes)


def value_on_or_before(series, target):
    prior = [v for d, v in series if d <= target]
    return prior[-1] if prior else None


def change(last, base, chg_type):
    if base is None or last is None or (chg_type != "bp" and base == 0):
        return ""
    if chg_type == "bp":
        return round((last - base) * 100, 1)
    return round((last / base - 1) * 100, 2)


def snapshot_row(u, series):
    last_d, last = series[-1]
    daily = u["frequency"] == "D"
    bases = {
        "chg_1d": series[-2][1] if daily and len(series) > 1 else None,
        "chg_1w": value_on_or_before(series, last_d - timedelta(days=7)) if daily else None,
        "chg_1m": value_on_or_before(series, last_d - timedelta(days=30)),
        "chg_3m": value_on_or_before(series, last_d - timedelta(days=91)),
        "chg_6m": value_on_or_before(series, last_d - timedelta(days=182)),
        "chg_ytd": value_on_or_before(series, date(last_d.year - 1, 12, 31)),
        "chg_1y": value_on_or_before(series, years_ago(last_d, 1)),
        "chg_2y": value_on_or_before(series, years_ago(last_d, 2)),
        "chg_3y": value_on_or_before(series, years_ago(last_d, 3)),
        "chg_5y": value_on_or_before(series, years_ago(last_d, 5)),
    }
    row = {k: u[k] for k in ["code", "name", "type", "region", "currency", "unit", "frequency", "chg_type", "headline"]}
    row["subtype"] = u.get("subtype", "")
    row.update({"as_of": last_d.isoformat(), "last": round(last, 4)})
    row.update({k: change(last, bases[k], u["chg_type"]) for k in HORIZONS})
    row.update(daily_stats(series, u["chg_type"]) if daily else {k: "" for k in STATS})
    row["source"] = SOURCE_LABEL.get(u["source"], u["source"])
    return row


def daily_stats(series, chg_type):
    """Moving averages, 52-week range and volatility signals for a daily series.
    sigma_1d   latest daily change / standard deviation of the 60 daily changes before it
    vol_20d    annualised volatility of the last 20 daily changes (percent, or bp for yields)
    signal_52w 'high' or 'low' when the latest close is the highest or lowest of the past year
    cross_200  'up' or 'down' when the close crossed the 200-day average in the last 3 trading days"""
    closes = [v for _, v in series]
    last_d, last = series[-1]
    bp = chg_type == "bp"
    out = {k: "" for k in STATS}

    def sma(vals, n):
        return sum(vals[-n:]) / n if len(vals) >= n else None

    ma50, ma200 = sma(closes, 50), sma(closes, 200)
    out.update({
        "ma50": round(ma50, 4) if ma50 else "", "ma200": round(ma200, 4) if ma200 else "",
        "vs_ma50": change(last, ma50, chg_type), "vs_ma200": change(last, ma200, chg_type),
    })
    year = [v for d, v in series if d > years_ago(last_d, 1)]
    hi, lo = max(year), min(year)
    out.update({
        "high_52w": round(hi, 4), "low_52w": round(lo, 4), "from_high_52w": change(last, hi, chg_type),
        "pos_52w": round((last - lo) / (hi - lo), 3) if hi > lo else "",
        "signal_52w": "high" if last >= hi else "low" if last <= lo else "",
    })
    moves = [(b - a) * 100 if bp else b / a - 1 for a, b in zip(closes, closes[1:])]
    if len(moves) >= 61:
        sd = statistics.stdev(moves[-61:-1])
        out["sigma_1d"] = round(moves[-1] / sd, 1) if sd else ""
    if len(moves) >= 20:
        recent = moves[-20:] if bp else [math.log(1 + m) for m in moves[-20:]]
        out["vol_20d"] = round(statistics.stdev(recent) * math.sqrt(252) * (1 if bp else 100), 1)
    if len(closes) >= 203:
        above = [closes[i] > sum(closes[i - 199:i + 1]) / 200 for i in range(len(closes) - 4, len(closes))]
        if above[-1] and not all(above[:-1]):
            out["cross_200"] = "up"
        elif not above[-1] and any(above[:-1]):
            out["cross_200"] = "down"
    return out


def read_csv(path):
    if not path.exists():
        return []
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def main():
    backfill = "--backfill" in sys.argv[1:]
    OUT.mkdir(exist_ok=True)
    universe = read_csv(ROOT / "universe.csv")
    today = datetime.now(SGT).date()
    begin = years_ago(today, YEARS) - timedelta(days=14)
    arc = load_archive()

    master, snapshot, failed, notes = [], [], [], []
    for u in universe:
        have = arc[(u["code"], u["field"], u["frequency"])]
        try:
            note = refresh(u, have, today, backfill)
            if note:
                notes.append(f"{u['code']}: {note}")
        except Exception as e:
            failed.append(f"{u['code']}: {e}" + (f" (kept {len(have)} archived rows, to {max(have)})" if have else ""))
        series = [(d, have[d][0]) for d in sorted(have) if d >= begin]
        master += [
            {
                "date": d.isoformat(), "frequency": u["frequency"], "code": u["code"], "name": u["name"],
                "type": u["type"], "field": u["field"], "value": round(v, 4), "unit": u["unit"],
                "currency": u["currency"], "source": SOURCE_LABEL.get(u["source"], u["source"]),
            }
            for d, v in series
        ]
        if series:
            snapshot.append(snapshot_row(u, series))

    if len(failed) == len(universe):
        sys.exit("Nothing fetched; files left unchanged.\n" + "\n".join(failed))

    save_archive(arc)
    write_csv(LONG, master, MASTER_COLS)
    write_csv(DASH, [{c: r[c] for c in DASH_COLS} for r in master], DASH_COLS)
    for freq, (fname, _) in WIDE_FILES.items():
        table = wide_table(master, freq)
        if table:
            with open(OUT / fname, "w", newline="", encoding="utf-8") as f:
                csv.writer(f).writerows(table)
    write_csv(OUT / "market_snapshot.csv", snapshot, list(snapshot[0].keys()))
    write_xlsx(OUT / "Market_Data.xlsx", snapshot, master, universe)

    stamp = datetime.now(SGT).strftime("%Y-%m-%d %H:%M SGT")
    print(f"{stamp}: {len(snapshot)} instruments, {len(master)} master rows, "
          f"{sum(map(len, arc.values()))} archive rows")
    for n in notes:
        print("NOTE", n)
    for f in failed:
        print("FAILED", f)

    try:
        import build_calendar
        build_calendar.main()
    except Exception as e:
        print("FAILED earnings calendar:", e)


def wide_table(master, freq):
    """Master-file layout for one frequency: metadata rows, a blank row, then date x instrument.
    A blank cell means the instrument has no value that day (e.g. its exchange was closed)."""
    rows = [r for r in master if r["frequency"] == freq]
    series = list(dict.fromkeys((r["code"], r["field"]) for r in rows))
    if not series:
        return None
    meta, values = {}, defaultdict(dict)
    for r in rows:
        k = (r["code"], r["field"])
        meta.setdefault(k, r)
        values[r["date"]][k] = r["value"]
    table = [[m] + [meta[k][m] for k in series] for m in META]
    table.append([])
    table.append(["date"] + [meta[k]["name"] for k in series])
    table += [[d] + [values[d].get(k, "") for k in series] for d in sorted(values)]
    return table


def write_csv(path, rows, cols):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)


def write_sheet(ws, cols, rows):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "B2"


def write_xlsx(path, snapshot, master, universe):
    wb = Workbook()
    ws = wb.active
    ws.title = "Snapshot"
    cols = list(snapshot[0].keys())
    write_sheet(ws, cols, [[r[c] for c in cols] for r in snapshot])

    # Master sheets in the master-file layout, with real Excel dates and numbers
    header_row = len(META) + 2  # 1-based row of the "date" header
    for freq, (_, sheet) in WIDE_FILES.items():
        table = wide_table(master, freq)
        if not table:
            continue
        ws = wb.create_sheet(sheet)
        for i, row in enumerate(table, start=1):
            if i > header_row:
                ws.append([date.fromisoformat(row[0])] + [float(v) if v != "" else None for v in row[1:]])
                ws.cell(i, 1).number_format = "dd/mm/yyyy"
            else:
                ws.append(row)
        for c in ws["A"][:header_row]:
            c.font = Font(bold=True)
        for c in ws[header_row]:
            c.font = Font(bold=True)
        ws.column_dimensions["A"].width = 12
        ws.freeze_panes = ws.cell(header_row + 1, 2)

    ucols = list(universe[0].keys())
    write_sheet(wb.create_sheet("Universe"), ucols, [[r[c] for c in ucols] for r in universe])
    wb.save(path)


if __name__ == "__main__":
    main()
